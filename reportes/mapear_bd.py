"""Diccionario de datos (solo lectura): tablas, columnas y filas de las bases accesibles."""
import os, pyodbc, pandas as pd
from openpyxl.styles import Font, PatternFill
LOGINS = [(os.environ['DB_USER'], os.environ['DB_PASS']), (os.environ['DBA_USER'], os.environ['DBA_PASS'])]
DBS = ['DBInforma', 'Analitica', 'Gestiones']
OUT = r'C:\Inoc\reportes\Mapa_BD_Informa.xlsx'

def conn(db, u, p):
    return pyodbc.connect(f"DRIVER={{ODBC Driver 18 for SQL Server}};SERVER=192.168.1.18;DATABASE={db};UID={u};PWD={p};"
                          "TrustServerCertificate=yes;Connection Timeout=10")
tablas, columnas, log = [], [], []
for db in DBS:
    c = None
    for u, p in LOGINS:
        try:
            c = conn(db, u, p); break
        except Exception:
            continue
    if c is None:
        log.append((db, 'sin acceso')); continue
    t = pd.read_sql("""SELECT s.name esquema, o.name tabla, o.type_desc tipo,
        SUM(CASE WHEN p.index_id IN (0,1) THEN p.rows ELSE 0 END) filas
        FROM sys.objects o JOIN sys.schemas s ON s.schema_id=o.schema_id
        LEFT JOIN sys.partitions p ON p.object_id=o.object_id
        WHERE o.type IN ('U','V') GROUP BY s.name,o.name,o.type_desc""", c)
    k = pd.read_sql("""SELECT TABLE_SCHEMA esquema, TABLE_NAME tabla, ORDINAL_POSITION pos, COLUMN_NAME columna, DATA_TYPE tipo_dato,
        CHARACTER_MAXIMUM_LENGTH largo, IS_NULLABLE nulo FROM INFORMATION_SCHEMA.COLUMNS""", c)
    t.insert(0, 'base', db); k.insert(0, 'base', db)
    tablas.append(t); columnas.append(k); log.append((db, f'{len(t)} objetos, {len(k)} columnas'))
    c.close()
T = pd.concat(tablas).sort_values(['base', 'esquema', 'tabla']); K = pd.concat(columnas)
ncol = K.groupby(['base', 'esquema', 'tabla']).columna.count().rename('n_columnas').reset_index()
cols = K.groupby(['base', 'esquema', 'tabla']).columna.apply(lambda s: ', '.join(s)).rename('columnas').reset_index()
T = T.merge(ncol, on=['base', 'esquema', 'tabla'], how='left').merge(cols, on=['base', 'esquema', 'tabla'], how='left')
with pd.ExcelWriter(OUT, engine='openpyxl') as xw:
    T.to_excel(xw, sheet_name='Tablas', index=False); K.to_excel(xw, sheet_name='Columnas', index=False)
    for ws in xw.sheets.values():
        for c in ws[1]:
            c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='1F4E78')
        ws.freeze_panes = 'A2'; ws.auto_filter.ref = ws.dimensions
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = min(60, max(10, max(len(str(x.value or '')) for x in col[:200]) + 2))
print(log); print(T.groupby('base').agg(objetos=('tabla', 'count'), filas=('filas', 'sum')))
