"""Reporte Cusco Extrajudicial - septiembre 2026, a partir de los archivos entregados.
Asignacion (16,767 cuentas): [Analitica].[Cusco].[ExtraJudicial_Asignacion] Periodo 202609  (solo SELECT)
Gestiones Sep-2026: contacto.csv | SMS e IVR Sep-2026: SMS_IVR_CUSCO.xlsx
Historico (cualquier mes): Base.Gestiones de DBInforma (solo SELECT). Ninguna cuenta se descarta."""
import os, re, pyodbc, pandas as pd
from openpyxl.styles import Font, PatternFill

DL = r'C:\Users\INFORMA PERU\Downloads'
OUT = r'C:\Inoc\reportes\Reporte_Cusco_Extrajudicial_Setiembre_2026_v7.xlsx'
CONTACTO = ['CONTACTO DIRECTO', 'CONTACTO INDIRECTO', 'CONTACTADO']
GCOLS = ['ID', 'FechaGestion', 'dni', 'telefono', 'TipoGestion', 'TipoContacto', 'TipoResultado', 'Observacion',
         'Fechapago', 'MontoPago', 'NombreAgente', 'Cartera', 'NroCuenta']


def conectar(db, u, p):
    return pyodbc.connect(f"DRIVER={{ODBC Driver 18 for SQL Server}};SERVER=192.168.1.18;DATABASE={db};UID={u};PWD={p};"
                          "TrustServerCertificate=yes;Connection Timeout=15")


# ---- Asignacion (universo) ----
ca = conectar('Analitica', os.environ['DBA_USER'], os.environ['DBA_PASS'])
b = pd.read_sql("""SELECT Cuenta, Doc_Titular, Saldo_Capital, Telefono
FROM [Analitica].[Cusco].[ExtraJudicial_Asignacion] WHERE Periodo = '202609'""", ca)
ca.close()
base = pd.DataFrame({'pagare': b.Cuenta.str.strip(), 'dni': b.Doc_Titular.str.strip().str.zfill(8),
                     'saldo': pd.to_numeric(b.Saldo_Capital, errors='coerce'), 'tels': b.Telefono.fillna('')})
assert base.pagare.is_unique

# ---- Gestiones septiembre: contacto.csv (sin encabezado, 13 columnas) ----
g = pd.read_csv(os.path.join(DL, 'contacto.csv'), dtype=str, encoding='utf-8-sig', sep=None, engine='python', header=None).iloc[:, :13]
g.columns = GCOLS
# NroCuenta puede traer 2 pagares del mismo cliente separados por salto de linea: se acredita a cada uno
g['pagare'] = g.NroCuenta.str.strip().str.split(r'\s+')
g = g.explode('pagare')
g['FechaGestion'] = pd.to_datetime(g.FechaGestion, errors='coerce')
g_all = g.pagare
g = g[g.pagare.isin(base.pagare)].copy()
g['contacto'] = g.TipoContacto.isin(CONTACTO)
ult_g = g.groupby('pagare').FechaGestion.max().rename('UltGest_sep')
ult_c = g[g.contacto].groupby('pagare').FechaGestion.max().rename('UltCont_sep')
llam = g.groupby('pagare').size().rename('Llamadas')
wa = g[g.Observacion.fillna('').str.upper().str.contains(r'WSP|WHATS|WAPP|\bWS\b', regex=True)].groupby('pagare').size().rename('WhatsApp')

# ---- SMS e IVR septiembre: SMS_IVR_CUSCO.xlsx ----
f = os.path.join(DL, 'SMS_IVR_CUSCO.xlsx')
sms = pd.read_excel(f, sheet_name='SMS', dtype=str)
ivr = pd.read_excel(f, sheet_name='IVR', dtype=str)
last9 = lambda t: re.sub(r'\D', '', str(t))[-9:]
numcol = [c for c in sms.columns if c.upper().startswith('N') and 'MERO' in c.upper()][0]
smsd = sms[numcol].map(last9).value_counts()
tb = base[['pagare', 'tels']].assign(tel=base.tels.str.split(r'[,;/ ]+')).explode('tel')
tg = g[['pagare', 'telefono']].rename(columns={'telefono': 'tel'})
t = pd.concat([tb[['pagare', 'tel']], tg]).dropna()
t['tel'] = t.tel.map(last9)
t = t[t.tel.str.len() == 9].drop_duplicates()
t['sms'] = t.tel.map(smsd).fillna(0)
sms_c = t.groupby('pagare').sms.sum().rename('SMS')
ivr_d = ivr.Documento.str.strip().str.zfill(8).value_counts()

# ---- Historico (cualquier mes) desde Base.Gestiones ----
cj = conectar('DBInforma', os.environ['DB_USER'], os.environ['DB_PASS'])
h = pd.read_sql("""SELECT NroCuenta pagare, FechaGestion, TipoContacto FROM Base.Gestiones
WHERE NroCuenta IS NOT NULL AND Cartera LIKE '%CUSCO%EXTRA%'""", cj)
cj.close()
h['pagare'] = h.pagare.str.strip().str.split(r'\s+')
h = h.explode('pagare')
h = h[h.pagare.isin(base.pagare)]
h_g = h.groupby('pagare').FechaGestion.max().rename('UltGest_hist')
h_c = h[h.TipoContacto.isin(CONTACTO)].groupby('pagare').FechaGestion.max().rename('UltCont_hist')

# ---- Armado ----
r = base.set_index('pagare').join([ult_c, ult_g, llam, sms_c, wa, h_c, h_g]).reset_index()
for c in ['Llamadas', 'SMS', 'WhatsApp']:
    r[c] = r[c].fillna(0).astype(int)
r['IVR'] = r.dni.map(ivr_d).fillna(0).astype(int)

out = r.rename(columns={'pagare': 'Pagaré', 'dni': 'DNI', 'saldo': 'Saldo Capital',
                        'UltCont_sep': 'Fecha Último Contacto (Sep-2026)', 'UltGest_sep': 'Fecha Última Gestión (Sep-2026)',
                        'SMS': 'Cant. SMS', 'IVR': 'Cant. IVR',
                        'UltCont_hist': 'Último Contacto Histórico', 'UltGest_hist': 'Última Gestión Histórica'})[
    ['Pagaré', 'DNI', 'Saldo Capital', 'Fecha Último Contacto (Sep-2026)', 'Fecha Última Gestión (Sep-2026)',
     'Llamadas', 'Cant. SMS', 'WhatsApp', 'Cant. IVR', 'Último Contacto Histórico', 'Última Gestión Histórica']]
for c in ['Fecha Último Contacto (Sep-2026)', 'Fecha Última Gestión (Sep-2026)', 'Último Contacto Histórico', 'Última Gestión Histórica']:
    out[c] = pd.to_datetime(out[c]).dt.strftime('%d/%m/%Y %H:%M')

# controles de integridad
assert len(out) == len(base) and out['Pagaré'].is_unique
assert abs(out['Saldo Capital'].sum() - base.saldo.sum()) < 0.01
assert out.Llamadas.sum() == len(g), 'las llamadas del reporte no suman las gestiones del CSV'

with pd.ExcelWriter(OUT, engine='openpyxl') as xw:
    out.to_excel(xw, index=False, sheet_name='Septiembre 2026')
    ws = xw.sheets['Septiembre 2026']
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor='1F4E78')
    for col, w in zip('ABCDEFGHIJK', [22, 12, 14, 24, 24, 10, 10, 10, 10, 24, 24]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    for row in ws.iter_rows(min_row=2, min_col=3, max_col=3):
        row[0].number_format = '#,##0.00'

n = lambda c: out[c].notna().sum()
print(f"{len(out)} filas = base ({len(base)}) | saldo {out['Saldo Capital'].sum():,.2f}")
print('gestiones CSV (por pagare):', len(g), '(CSV: 1311 filas; 47 traen 2 pagares) | cuentas con llamadas:', (out.Llamadas > 0).sum(), '| suma llamadas:', out.Llamadas.sum())
print('ult.contacto Sep:', n('Fecha Último Contacto (Sep-2026)'), '| SMS:', (out['Cant. SMS'] > 0).sum(), '| WhatsApp:', (out.WhatsApp > 0).sum(),
      '| IVR:', (out['Cant. IVR'] > 0).sum())
print('Hist: ult.contacto', n('Último Contacto Histórico'), '| ult.gestion', n('Última Gestión Histórica'))
print('gestiones del CSV con pagare fuera de la base:', (~g_all.isin(base.pagare)).sum(), '(g_all = pagares del CSV tras separar)')
